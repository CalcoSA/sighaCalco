from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from unicodedata import normalize
from io import BytesIO
import re

class ReconciliationExcelReader:

    DOCUMENT_HEADERS = {
        "documento",
        "cedula",
        "identificacion",
        "numero documento",
        "numero de documento",
        "documento colaborador",
    }

    NAME_HEADERS = {
        "nombre",
        "nombre completo",
        "colaborador",
        "empleado",
    }

    AMOUNT_HEADERS = {
        "cuota",
        "valor cuota",
        "valor de cuota",
        "valor",
        "descuento",
        "valor descuento",
    }

    def read(self, content: bytes,) -> list[dict]:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True,)
        worksheet = self._findWorksheet(workbook)

        if worksheet is None:
            raise ValueError(
                "No fue posible encontrar una hoja con las columnas Documento, Nombre y Cuota.")

        (headerRow, documentColumn, nameColumn, amountColumn,) = self._findHeaders(worksheet)
        records: list[dict] = []

        for rowNumber in range(headerRow + 1, worksheet.max_row + 1,):
            documentValue = worksheet.cell(rowNumber, documentColumn,).value
            nameValue = worksheet.cell(rowNumber, nameColumn,).value
            amountValue = worksheet.cell(rowNumber, amountColumn,).value

            if (documentValue is None and nameValue is None and amountValue is None):
                continue

            documentNumber = (self._normalizeDocument(documentValue))

            if not documentNumber:
                raise ValueError(f"La fila {rowNumber} no tiene un documento válido.")

            amount = self._parseAmount(amountValue)

            if amount is None:
                raise ValueError(f"La fila {rowNumber} no tiene una cuota válida.")

            records.append(
                {
                    "rowNumber": rowNumber,
                    "documentNumber": documentNumber,
                    "fullName": str(nameValue or "").strip(),
                    "amount": amount,
                }
            )

        return records

    def _findWorksheet(self, workbook,):
        for worksheet in workbook.worksheets:
            try:
                self._findHeaders(worksheet)
                return worksheet

            except ValueError:
                continue

        return None

    def _findHeaders(self, worksheet,):
        maxRows = min(worksheet.max_row, 20,)

        for rowNumber in range(1, maxRows + 1,):
            documentColumn = None
            nameColumn = None
            amountColumn = None

            for columnNumber in range(1, worksheet.max_column + 1,):
                value = worksheet.cell(rowNumber, columnNumber,).value
                header = (self._normalizeHeader(value))

                if (header in self.DOCUMENT_HEADERS):
                    documentColumn = (columnNumber)
                elif (header in self.NAME_HEADERS):
                    nameColumn = (columnNumber)
                elif (header in self.AMOUNT_HEADERS):
                    amountColumn = (columnNumber)

            if (documentColumn and nameColumn and amountColumn):
                return (rowNumber, documentColumn, nameColumn, amountColumn,)

        raise ValueError("No se encontraron los encabezados.")

    def _normalizeHeader(self, value,) -> str:

        if value is None:
            return ""

        text = (str(value).strip().lower())
        text = "".join(character
            for character in normalize("NFD", text,)
            if ord(character) < 128
        )

        return re.sub(r"\s+", " ", text,)

    def _normalizeDocument(self, value,) -> str:

        if value is None:
            return ""

        if (isinstance(value, float) and value.is_integer()):
            value = int(value)

        return re.sub(r"\D", "", str(value),)

    def _parseAmount(self, value,) -> Decimal | None:

        if value is None:
            return None

        if isinstance(value, (int, float, Decimal),):
            return Decimal(str(value)).quantize(Decimal("0.01"))

        text = (str(value).strip().replace("$", "").replace(" ", ""))

        if not text:
            return None

        if ("." in text and "," in text):
            if (text.rfind(",") > text.rfind(".")):
                text = (text.replace(".", "").replace(",", "."))
            else:
                text = text.replace(",", "",)

        elif "," in text:
            text = text.replace(",", ".",)

        try:
            return Decimal(text).quantize(Decimal("0.01"))

        except InvalidOperation:
            return None